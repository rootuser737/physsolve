import os, subprocess, sys, tempfile
os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import numpy as np
import exporter, speech, teacher
from systems import SYSTEMS, METHODS, reference, run_method
from theme import THEMES, mpl_colors
from PySide6.QtGui import QGuiApplication


def make_state(system=None, T=6, dt=0.02):
    sysm = system or SYSTEMS[0]
    p = {k: d for k, _, d, *_ in sysm.params}
    t = np.linspace(0, int(round(T / dt)) * dt, int(round(T / dt)) + 1)
    Yref = reference(sysm, p, t)
    res = {m: run_method(sysm, p, m, dt, T, Yref) for m in METHODS}
    r = res[METHODS[1]]
    P = np.array([sysm.xy(y, p) for y in r.Y]); Pref = np.array([sysm.xy(y, p) for y in Yref])
    return dict(system=sysm, p=p, results=res, sel=METHODS[1], Yref=Yref, t=t, P=P, Pref=Pref, dt=dt, T=T,
                events=teacher.detect_events(sysm, p, Yref, t), hold=20)


def probe(path):
    import imageio_ffmpeg
    return imageio_ffmpeg.count_frames_and_secs(path)


def test_all():
    app = QGuiApplication.instance() or QGuiApplication([])
    d = tempfile.mkdtemp(); P = lambda n: os.path.join(d, n)
    colors = mpl_colors(THEMES["Ocean"])
    st = make_state()

    # ---- images
    from PIL import Image
    exporter.save_figure(P("a.png"), st, colors); im = Image.open(P("a.png")); assert im.size == (1600, 900), im.size
    exporter.save_figure(P("a.svg"), st, colors); assert "<svg" in open(P("a.svg"), encoding="utf-8").read(2000)
    exporter.save_figure(P("a.pdf"), st, colors); assert open(P("a.pdf"), "rb").read(5) == b"%PDF-"
    print("images ok")

    # ---- video: mp4 (frames/duration), gif, cancel, narration muxing
    prog = []
    note = exporter.render_video(st, colors, P("v.mp4"), size=(640, 360), fps=10, seconds=2, progress=prog.append)
    n, secs = probe(P("v.mp4")); assert n == 20 and 1.9 < secs < 2.2 and prog[-1] == 1.0, (n, secs, note)
    exporter.render_video(st, colors, P("v.gif"), size=(640, 360), fps=10, seconds=1)
    assert open(P("v.gif"), "rb").read(3) == b"GIF"
    try:
        exporter.render_video(st, colors, P("c.mp4"), size=(320, 180), fps=10, seconds=2, cancelled=lambda: True); assert False
    except exporter.Cancelled:
        pass
    ff = exporter.find_ffmpeg()
    def fake_tts(text, path, rate=0):
        subprocess.run([ff, "-y", "-loglevel", "error", "-f", "lavfi", "-i", "sine=f=300:duration=3", path], check=True); return True
    real = speech.synth_to_wav; speech.synth_to_wav = fake_tts
    exporter.render_video(st, colors, P("n.mp4"), size=(640, 360), fps=10, seconds=2, narration="Hello class")
    speech.synth_to_wav = real
    info = subprocess.run([ff, "-i", P("n.mp4")], capture_output=True, text=True).stderr
    assert "Audio:" in info and probe(P("n.mp4"))[1] >= 3.8, info[-300:]
    # when speech is unavailable the video is still produced, with an explanatory note
    speech.synth_to_wav = lambda *a, **k: False
    note = exporter.render_video(st, colors, P("n2.mp4"), size=(320, 180), fps=10, seconds=1, narration="x"); speech.synth_to_wav = real
    assert "no sound" in note
    print("video ok (mp4 frames=%d, gif, cancel, narration mux, fallback)" % n)

    # ---- documents: re-open every file and check real content
    t, Y = st["t"], st["Yref"]
    expl = teacher.explain(st["system"], st["p"], st["results"], st["sel"], st["dt"], st["T"], Y, t, "Undergraduate")
    expl += "\n\n" + teacher.convergence_text("RK4 (fixed step)", teacher.convergence_study(
        st["system"], st["p"], "RK4 (fixed step)", 0.1, 4, lambda tt: reference(st["system"], st["p"], tt)))
    chat = [{"role": "user", "content": "Why does Euler drift?"}, {"role": "assistant", "content": "Because $E=mc^2$ is **not** kept by Euler.\n\n- one\n- two"},
            {"role": "system", "content": "internal note"}]
    rep = exporter.make_report(st, expl, chat, True, colors, "My lab note")
    assert len(rep["chat"]) == 2 and len(rep["images"]) == 2 and len(rep["results"]) == 3
    exporter.save_docx(rep, P("r.docx"))
    import docx
    doc = docx.Document(P("r.docx"))
    assert len(doc.tables) >= 3 and len(doc.inline_shapes) == 2 and any("Lab report" in p.text for p in doc.paragraphs)
    assert any("Tutor:" in p.text for p in doc.paragraphs) and not any("$" in p.text for p in doc.paragraphs)
    exporter.save_pdf(rep, P("r.pdf")); b = open(P("r.pdf"), "rb").read()
    assert b[:5] == b"%PDF-" and len(b) > 30000 and b.count(b"/Type /Page") >= 1, len(b)
    exporter.save_html(rep, P("r.html")); h = open(P("r.html"), encoding="utf-8").read()
    assert h.count("data:image/png;base64") == 2 and "<table>" in h and "$" not in h and "Convergence" in h
    exporter.save_markdown(rep, P("r.md")); md = open(P("r.md"), encoding="utf-8").read()
    assert "![Simulation dashboard](r_images/figure1.png)" in md and os.path.getsize(P("r_images/figure2.png")) > 5000
    print("documents ok (docx %d KB, pdf %d KB, html %d KB)" % (os.path.getsize(P("r.docx")) // 1000, len(b) // 1000, len(h) // 1000))

    # ---- all four systems export cleanly (different shapes of P / Pref)
    for s in SYSTEMS[1:]:
        s2 = make_state(s, T=3, dt=0.01); exporter.save_figure(P(f"{s.name[:3]}.png"), s2, colors, size=(800, 450))
    # ---- session round trip
    exporter.save_session(P("s.sciproj"), {"system": "Kepler", "params": {"e": 0.9}, "chat": chat})
    s = exporter.load_session(P("s.sciproj")); assert s["params"]["e"] == 0.9 and len(s["chat"]) == 3
    open(P("bad.json"), "w").write('{"x":1}')
    try:
        exporter.load_session(P("bad.json")); assert False
    except ValueError:
        pass
    print("EXPORT OK")


if __name__ == "__main__":
    test_all()
