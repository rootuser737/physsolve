"""Saving: images (PNG/SVG/PDF), videos (MP4/GIF, optional spoken narration), documents (DOCX/PDF/HTML/Markdown), sessions."""
import base64
import datetime
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from html import escape
from pathlib import Path

import numpy as np
from matplotlib.backends.backend_agg import FigureCanvasAgg
from matplotlib.figure import Figure

import speech
from plots import LabPlots

NOWIN = subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0


# ------------------------------------------------------------------ locations
def notebook_dir() -> Path:
    for base in (Path.home() / "Documents", Path.home()):
        try:
            d = base / "SciSketch"
            d.mkdir(parents=True, exist_ok=True)
            return d
        except OSError:
            continue
    return Path(tempfile.gettempdir())


def default_name(label: str, ext: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9]+", "_", label).strip("_")[:28] or "lab"
    return str(notebook_dir() / f"{slug}_{datetime.datetime.now():%Y%m%d_%H%M%S}.{ext}")


def find_ffmpeg():
    exe = "ffmpeg.exe" if os.name == "nt" else "ffmpeg"
    for c in (Path(os.sys.executable).parent / exe, Path(__file__).parent / "bin" / exe):
        if c.is_file():
            return str(c)
    try:
        import imageio_ffmpeg
        return imageio_ffmpeg.get_ffmpeg_exe()
    except Exception:  # noqa: BLE001
        return shutil.which("ffmpeg")


# ------------------------------------------------------------------ figures & images
def make_fig(state, colors, size=(1280, 720), dpi=100):
    fig = Figure(figsize=(size[0] / dpi, size[1] / dpi), dpi=dpi, constrained_layout=True)
    FigureCanvasAgg(fig)
    lp = LabPlots(fig, colors)
    lp.build(state["system"], state["p"], state["results"], state["sel"], state["Yref"], state["t"], state["P"], state["Pref"])
    lp.set_events(state.get("events", []), state.get("hold", 1))
    return fig, lp


def frame_default(state):
    return int(len(state["t"]) * 0.35)


def save_figure(path, state, colors, frame=None, size=(1600, 900)):
    """PNG / SVG / PDF of the whole dashboard at the chosen frame."""
    fig, lp = make_fig(state, colors, size, dpi=100)
    lp.update(frame if frame is not None else frame_default(state))
    fig.savefig(path, dpi=100 if path.lower().endswith(".png") else None, facecolor=fig.get_facecolor())
    return path


def png_bytes(fig):
    b = io.BytesIO()
    fig.savefig(b, format="png", dpi=fig.dpi, facecolor=fig.get_facecolor())
    return b.getvalue()


def energy_figure(state, colors):
    from plots import MCOL
    fig = Figure(figsize=(7.2, 3.2), dpi=110, constrained_layout=True, facecolor=colors["fig"])
    FigureCanvasAgg(fig)
    ax = fig.add_subplot(111)
    ax.set_facecolor(colors["ax"]); ax.set_yscale("log"); ax.set_xlabel("t (s)", color=colors["muted"])
    ax.set_title("Energy drift  |ΔE| / max(|E₀|,1)", color=colors["text"], fontsize=10)
    ax.tick_params(colors=colors["muted"], labelsize=8); ax.grid(color=colors["grid"], alpha=.7, lw=.6)
    for s in ax.spines.values():
        s.set_color(colors["grid"])
    for m, r in state["results"].items():
        ax.plot(r.t, np.maximum(r.drift, 1e-17), color=MCOL[m], label=m, lw=1.5)
    ax.axhline(1e-6, color="#3FD08A", ls=":", lw=1); ax.axhline(1e-2, color="#F0506E", ls=":", lw=1)
    ax.set_ylim(1e-17, 1e3)
    lg = ax.legend(fontsize=7, facecolor=colors["ax"], edgecolor=colors["grid"])
    for tx in lg.get_texts():
        tx.set_color(colors["text"])
    return fig


# ------------------------------------------------------------------ video
class Cancelled(Exception):
    pass


def render_video(state, colors, path, size=(1280, 720), fps=30, seconds=20, narration=None,
                 progress=lambda f: None, cancelled=lambda: False):
    """Render the whole simulation (all panels, live captions) to MP4 or GIF. Returns a status note."""
    ff = find_ffmpeg()
    if not ff:
        raise RuntimeError("FFmpeg was not found. Reinstall SciSketch or put ffmpeg next to the app.")
    w, h = size[0] - size[0] % 2, size[1] - size[1] % 2
    gif = path.lower().endswith(".gif")
    audio, note = None, ""
    if narration and not gif:
        wav = os.path.join(tempfile.gettempdir(), f"scisketch_nar_{os.getpid()}.wav")
        if speech.synth_to_wav(narration, wav):
            dur = _duration(ff, wav)
            seconds, audio = max(seconds, dur + 1.0), wav
        else:
            note = " (voice narration is not available on this computer, so the video has no sound)"
    n = max(2, int(seconds * fps))
    idx = np.linspace(0, len(state["t"]) - 1, n).astype(int)
    fig, lp = make_fig(state, colors, (w, h), dpi=100)
    lp.set_events(state.get("events", []), max(1, int(len(state["t"]) * 0.06)))
    base = [ff, "-y", "-loglevel", "error", "-f", "rawvideo", "-pix_fmt", "rgba", "-s", f"{w}x{h}", "-r", str(fps), "-i", "-"]
    if gif:
        gw = min(w, 900)
        cmd = base + ["-vf", f"fps={min(fps, 15)},scale={gw}:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse", path]
    else:
        cmd = base + (["-i", audio] if audio else []) + ["-c:v", "libx264", "-preset", "veryfast", "-crf", "18",
                                                         "-pix_fmt", "yuv420p", "-movflags", "+faststart"]
        cmd += (["-c:a", "aac", "-b:a", "160k"] if audio else []) + [path]
    err = tempfile.TemporaryFile()
    proc = subprocess.Popen(cmd, stdin=subprocess.PIPE, stderr=err, creationflags=NOWIN)
    try:
        for k, i in enumerate(idx):
            if cancelled():
                raise Cancelled()
            lp.update(int(i))
            fig.canvas.draw()
            proc.stdin.write(fig.canvas.buffer_rgba().tobytes())
            if k % 4 == 0:
                progress(k / n)
        proc.stdin.close()
        proc.wait()
    except Cancelled:
        proc.kill()
        raise
    except (BrokenPipeError, OSError):
        err.seek(0)
        raise RuntimeError("FFmpeg stopped unexpectedly:\n" + err.read().decode("utf-8", "replace")[-500:])
    finally:
        if audio and os.path.exists(audio):
            os.remove(audio)
    if proc.returncode != 0:
        err.seek(0)
        raise RuntimeError("FFmpeg failed:\n" + err.read().decode("utf-8", "replace")[-500:])
    progress(1.0)
    return f"{n} frames, {seconds:.0f} s{note}"


def _duration(ff, path):
    r = subprocess.run([ff, "-i", path], capture_output=True, text=True, creationflags=NOWIN)
    m = re.search(r"Duration: (\d+):(\d+):([\d.]+)", r.stderr)
    return int(m[1]) * 3600 + int(m[2]) * 60 + float(m[3]) if m else 10.0


# ------------------------------------------------------------------ documents
_TEX = [(r"\\frac\{([^{}]*)\}\{([^{}]*)\}", r"(\1)/(\2)"), (r"\\sqrt\{([^{}]*)\}", r"√(\1)"), (r"\\text\{([^{}]*)\}", r"\1"),
        (r"\\tfrac", ""), (r"\\left|\\right|\\!|\\,", ""), (r"\\pi", "π"), (r"\\theta", "θ"), (r"\\omega", "ω"),
        (r"\\Delta", "Δ"), (r"\\delta", "δ"), (r"\\lambda", "λ"), (r"\\approx", "≈"), (r"\\cdot", "·"), (r"\\sim", "~"),
        (r"\\to", "→"), (r"\\dot A", "Ȧ"), (r"\\sin", "sin"), (r"\\cos", "cos"), (r"\\in", "∈"), (r"\^2", "²"), (r"\^3", "³"),
        (r"\^\{([^{}]*)\}", r"^(\1)"), (r"_\{([^{}]*)\}", r"_\1"), (r"[{}]", ""), (r"\\", "")]


def tex_plain(md: str) -> str:
    def conv(m):
        s = m.group(1)
        for a, b in _TEX:
            s = re.sub(a, b, s)
        return s.strip()
    return re.sub(r"\$\$?(.+?)\$\$?", conv, md, flags=re.S)


def md_blocks(md: str):
    blocks, lines, i = [], tex_plain(md).splitlines(), 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1; continue
        m = re.match(r"(#{1,4})\s+(.*)", ln)
        if m:
            blocks.append(("h", len(m[1]), m[2])); i += 1
        elif ln.lstrip().startswith("|"):
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            blocks.append(("table", rows))
        elif re.match(r"\s*[-*]\s+", ln):
            items = []
            while i < len(lines) and re.match(r"\s*[-*]\s+", lines[i]):
                items.append(re.sub(r"^\s*[-*]\s+", "", lines[i])); i += 1
            blocks.append(("ul", items))
        else:
            para = [ln]; i += 1
            while i < len(lines) and lines[i].strip() and not re.match(r"(#{1,4}\s|\s*[-*]\s|\s*\|)", lines[i]):
                para.append(lines[i].rstrip()); i += 1
            blocks.append(("p", " ".join(para)))
    return blocks


def inline_html(s):
    s = escape(s)
    s = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"`([^`]+)`", r"<code>\1</code>", s)
    return re.sub(r"(?<![*\w])\*([^*\n]+)\*(?!\w)", r"<i>\1</i>", s)


def make_report(state, explanation, transcript, include_chat, colors, extra_notes=""):
    """Collect everything a lab report needs."""
    fig, lp = make_fig(state, colors, (1500, 800), dpi=100)
    lp.update(frame_default(state))
    imgs = [("Simulation dashboard", png_bytes(fig)), ("Energy conservation check", png_bytes(energy_figure(state, colors)))]
    res = []
    for m, r in state["results"].items():
        me = np.nanmax(r.err) if not np.all(np.isnan(r.err)) else float("nan")
        res.append([m, f"{r.max_drift:.3e}", f"{me:.3e}", f"{r.runtime * 1000:.0f} ms", r.verdict])
    params = [[k, f"{v:g}"] for k, v in state["p"].items()] + [["time step dt", f"{state['dt']:g} s"], ["duration T", f"{state['T']:g} s"],
                                                              ["displayed method", state["sel"]]]
    chat = [m for m in transcript if m["role"] in ("user", "assistant") and m["content"].strip()] if include_chat else []
    return dict(title=f"Lab report - {state['system'].name}", date=f"{datetime.datetime.now():%d %B %Y, %H:%M}", params=params,
                results=res, explanation=explanation, chat=chat, images=imgs, notes=extra_notes)


def save_markdown(rep, path):
    path = Path(path); assets = path.with_suffix("").with_name(path.stem + "_images"); assets.mkdir(exist_ok=True)
    out = [f"# {rep['title']}", f"*{rep['date']} - SciSketch*", "", "## Setup", "| Parameter | Value |", "|---|---|"]
    out += [f"| {a} | {b} |" for a, b in rep["params"]]
    out += ["", "## Verification against the DOP853 reference (rtol 1e-12)", "| Method | Max energy drift | Max state error | Runtime | Verdict |", "|---|---|---|---|---|"]
    out += ["| " + " | ".join(r) + " |" for r in rep["results"]] + [""]
    for k, (cap, data) in enumerate(rep["images"]):
        fn = f"{assets.name}/figure{k + 1}.png"
        (assets / f"figure{k + 1}.png").write_bytes(data)
        out += [f"![{cap}]({fn})", f"*Figure {k + 1}: {cap}*", ""]
    out += ["## Explanation", rep["explanation"], ""]
    if rep["notes"]:
        out += ["## Notes", rep["notes"], ""]
    if rep["chat"]:
        out += ["## Tutor conversation"] + [f"**{'You' if m['role'] == 'user' else 'Tutor'}:** {m['content']}\n" for m in rep["chat"]]
    path.write_text("\n".join(out), encoding="utf-8")
    return str(path)


_CSS = ("body{font-family:Segoe UI,Arial,sans-serif;max-width:900px;margin:30px auto;color:#1b1f3b;line-height:1.55;padding:0 16px}"
        "h1{color:#2b4acb}h2{border-bottom:2px solid #dde3f5;padding-bottom:4px;margin-top:28px}table{border-collapse:collapse;margin:8px 0}"
        "td,th{border:1px solid #cfd6ea;padding:5px 10px;font-size:14px}th{background:#eef1fb}img{max-width:100%;border:1px solid #dde3f5;border-radius:6px}"
        "code{background:#eef1fb;padding:1px 4px;border-radius:4px}.cap{color:#6a7096;font-size:13px;margin:2px 0 14px}"
        ".chat{background:#f4f6fc;border-radius:8px;padding:6px 12px;margin:6px 0}")


def _blocks_html(md):
    out = []
    for b in md_blocks(md):
        if b[0] == "h":
            out.append(f"<h{min(b[1] + 1, 4)}>{inline_html(b[2])}</h{min(b[1] + 1, 4)}>")
        elif b[0] == "p":
            out.append(f"<p>{inline_html(b[1])}</p>")
        elif b[0] == "ul":
            out.append("<ul>" + "".join(f"<li>{inline_html(x)}</li>" for x in b[1]) + "</ul>")
        else:
            out.append("<table>" + "".join("<tr>" + "".join(f"<{'th' if r == 0 else 'td'}>{inline_html(c)}</{'th' if r == 0 else 'td'}>" for c in row)
                                           + "</tr>" for r, row in enumerate(b[1])) + "</table>")
    return "\n".join(out)


def build_html(rep, embed=True, images=None):
    def table(head, rows):
        return ("<table><tr>" + "".join(f"<th>{escape(h)}</th>" for h in head) + "</tr>"
                + "".join("<tr>" + "".join(f"<td>{escape(c)}</td>" for c in r) + "</tr>" for r in rows) + "</table>")
    h = [f"<h1>{escape(rep['title'])}</h1><div class='cap'>{escape(rep['date'])} - SciSketch</div>", "<h2>Setup</h2>",
         table(["Parameter", "Value"], rep["params"]), "<h2>Verification against the DOP853 reference</h2>",
         table(["Method", "Max energy drift", "Max state error", "Runtime", "Verdict"], rep["results"])]
    for k, (cap, data) in enumerate(rep["images"]):
        src = f"data:image/png;base64,{base64.b64encode(data).decode()}" if embed else f"img{k}"
        h.append(f"<p><img src='{src}' width='{760}'></p><div class='cap'>Figure {k + 1}: {escape(cap)}</div>")
    h += ["<h2>Explanation</h2>", _blocks_html(rep["explanation"])]
    if rep["notes"]:
        h += ["<h2>Notes</h2>", f"<p>{inline_html(rep['notes'])}</p>"]
    if rep["chat"]:
        h.append("<h2>Tutor conversation</h2>")
        h += [f"<div class='chat'><b>{'You' if m['role'] == 'user' else 'Tutor'}:</b> {_blocks_html(m['content'])}</div>" for m in rep["chat"]]
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{escape(rep['title'])}</title><style>{_CSS}</style></head><body>{''.join(h)}</body></html>"


def save_html(rep, path):
    Path(path).write_text(build_html(rep), encoding="utf-8")
    return str(path)


def save_pdf(rep, path):
    from PySide6.QtCore import QMarginsF, QUrl
    from PySide6.QtGui import QImage, QPageLayout, QPageSize, QPdfWriter, QTextDocument
    w = QPdfWriter(str(path)); w.setPageSize(QPageSize(QPageSize.A4)); w.setResolution(96)
    w.setPageMargins(QMarginsF(14, 14, 14, 14), QPageLayout.Millimeter)
    doc = QTextDocument()
    doc.setDefaultStyleSheet("table{border-collapse:collapse} td,th{border:1px solid #bbb;padding:4px} th{background:#e9edf9} "
                             "h1{color:#2b4acb} h2{color:#2b4acb} code{background:#eef1fb}")
    for k, (_, data) in enumerate(rep["images"]):
        doc.addResource(QTextDocument.ImageResource, QUrl(f"img{k}"), QImage.fromData(data))
    html = build_html(rep, embed=False).replace("width='760'", "width='680'")
    doc.setHtml(html)
    doc.print_(w)
    return str(path)


def save_docx(rep, path):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor

    d = Document()
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Calibri", Pt(11)

    def runs(par, text):
        for tok in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
            if tok.startswith("**"):
                par.add_run(tok[2:-2]).bold = True
            elif tok.startswith("`"):
                r = par.add_run(tok[1:-1]); r.font.name = "Consolas"
            elif tok:
                par.add_run(tok.replace("*", ""))

    def table(head, rows):
        t = d.add_table(rows=1, cols=len(head)); t.style = "Table Grid"
        for i, h in enumerate(head):
            c = t.rows[0].cells[i]; c.text = ""; c.paragraphs[0].add_run(h).bold = True
        for r in rows:
            cells = t.add_row().cells
            for i, v in enumerate(r):
                cells[i].text = ""; runs(cells[i].paragraphs[0], str(v))
        d.add_paragraph()

    def blocks(md):
        for b in md_blocks(md):
            if b[0] == "h":
                d.add_heading(re.sub(r"[*`]", "", b[2]), level=min(b[1] + 1, 4))
            elif b[0] == "p":
                runs(d.add_paragraph(), b[1])
            elif b[0] == "ul":
                for x in b[1]:
                    runs(d.add_paragraph(style="List Bullet"), x)
            else:
                table(b[1][0], b[1][1:])

    d.add_heading(rep["title"], 0)
    p = d.add_paragraph(); r = p.add_run(f"{rep['date']} - SciSketch"); r.italic = True; r.font.color.rgb = RGBColor(0x6A, 0x70, 0x96)
    d.add_heading("Setup", 1); table(["Parameter", "Value"], rep["params"])
    d.add_heading("Verification against the DOP853 reference", 1)
    table(["Method", "Max energy drift", "Max state error", "Runtime", "Verdict"], rep["results"])
    for k, (cap, data) in enumerate(rep["images"]):
        d.add_picture(io.BytesIO(data), width=Inches(6.3))
        c = d.add_paragraph(); c.add_run(f"Figure {k + 1}: {cap}").italic = True
    d.add_heading("Explanation", 1); blocks(rep["explanation"])
    if rep["notes"]:
        d.add_heading("Notes", 1); runs(d.add_paragraph(), rep["notes"])
    if rep["chat"]:
        d.add_heading("Tutor conversation", 1)
        for m in rep["chat"]:
            p = d.add_paragraph(); p.add_run("You: " if m["role"] == "user" else "Tutor: ").bold = True
            runs(p, re.sub(r"\s*\n\s*", " ", tex_plain(m["content"])))
    d.save(str(path))
    return str(path)


SAVERS = {"docx": save_docx, "pdf": save_pdf, "html": save_html, "md": save_markdown}


# ------------------------------------------------------------------ sessions
def save_session(path, data):
    data = dict(data, app="SciSketch", saved=datetime.datetime.now().isoformat(timespec="seconds"))
    Path(path).write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    return str(path)


def load_session(path):
    d = json.loads(Path(path).read_text(encoding="utf-8"))
    if d.get("app") != "SciSketch":
        raise ValueError("This is not a SciSketch session file.")
    return d
